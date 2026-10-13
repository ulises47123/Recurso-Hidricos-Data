import os
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from backend.database import get_db_connection
from backend.services.auditoria_service import registrar_auditoria

def exportar_libro_completo():
    """Genera un archivo Excel (.xlsx) con las 4 pestañas: Personal, Licencias, Contratos, Resumen"""
    wb = openpyxl.Workbook()
    
    # Colores corporativos Recursos Hídricos
    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid") # Azul Marino
    sub_fill = PatternFill(start_color="2B6CB0", end_color="2B6CB0", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Arial", size=10, bold=True)
    regular_font = Font(name="Arial", size=10)
    title_font = Font(name="Arial", size=14, bold=True, color="1A365D")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E0'),
        right=Side(style='thin', color='CBD5E0'),
        top=Side(style='thin', color='CBD5E0'),
        bottom=Side(style='thin', color='CBD5E0')
    )

    conn = get_db_connection()
    cursor = conn.cursor()

    # -------------------------------------------------------------
    # 1. PESTAÑA RESUMEN (DASHBOARD EXECUTIVO)
    # -------------------------------------------------------------
    ws_resumen = wb.active
    ws_resumen.title = "Resumen"
    ws_resumen.views.sheetView[0].showGridLines = True

    ws_resumen["A1"] = "SISTEMA DE GESTIÓN DE PERSONAL - RECURSOS HÍDRICOS"
    ws_resumen["A1"].font = title_font
    ws_resumen["A2"] = f"Reporte Consolidado generado el: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
    ws_resumen["A2"].font = Font(italic=True, size=10, color="4A5568")

    # Estadísticas Generales
    cursor.execute("SELECT COUNT(*) FROM personal WHERE deleted_at IS NULL")
    total_personal = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM personal WHERE estado = 'Activo' AND deleted_at IS NULL")
    activos = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM licencias WHERE deleted_at IS NULL AND estado_justificacion = 'Pendiente'")
    licencias_pendientes = cursor.fetchone()[0]

    cursor.execute("SELECT COUNT(*) FROM licencias WHERE deleted_at IS NULL")
    total_licencias = cursor.fetchone()[0]

    ws_resumen["A4"] = "INDICADOR GENERAL"
    ws_resumen["B4"] = "CANTIDAD"
    ws_resumen["A4"].fill = header_fill
    ws_resumen["A4"].font = header_font
    ws_resumen["B4"].fill = header_fill
    ws_resumen["B4"].font = header_font

    kpis = [
        ("Total de Empleados Registrados", total_personal),
        ("Personal en Estado Activo", activos),
        ("Licencias Totales Registradas", total_licencias),
        ("Certificados Médicos Pendientes de Justificación", licencias_pendientes),
    ]
    for idx, (label, val) in enumerate(kpis, start=5):
        ws_resumen[f"A{idx}"] = label
        ws_resumen[f"B{idx}"] = val
        ws_resumen[f"A{idx}"].font = regular_font
        ws_resumen[f"B{idx}"].font = bold_font
        ws_resumen[f"A{idx}"].border = thin_border
        ws_resumen[f"B{idx}"].border = thin_border

    # Distribución por Situación de Revista
    ws_resumen["D4"] = "SITUACIÓN DE REVISTA"
    ws_resumen["E4"] = "CANTIDAD"
    ws_resumen["D4"].fill = sub_fill
    ws_resumen["D4"].font = header_font
    ws_resumen["E4"].fill = sub_fill
    ws_resumen["E4"].font = header_font

    cursor.execute("""
        SELECT situacion_revista, COUNT(*) 
        FROM personal 
        WHERE deleted_at IS NULL 
        GROUP BY situacion_revista
    """)
    revistas = cursor.fetchall()
    for idx, row in enumerate(revistas, start=5):
        ws_resumen[f"D{idx}"] = row[0]
        ws_resumen[f"E{idx}"] = row[1]
        ws_resumen[f"D{idx}"].font = regular_font
        ws_resumen[f"E{idx}"].font = bold_font
        ws_resumen[f"D{idx}"].border = thin_border
        ws_resumen[f"E{idx}"].border = thin_border

    # -------------------------------------------------------------
    # 2. PESTAÑA PERSONAL
    # -------------------------------------------------------------
    ws_personal = wb.create_sheet(title="Personal")
    headers_p = [
        "ID", "Legajo", "CUIL", "DNI", "Apellido", "Nombre", "Fecha Nacimiento", "Sexo", 
        "Domicilio", "Teléfono", "Email", "Horario Trabajo", 
        "F. Ingreso Admin Pública", "F. Ingreso DPRH Obra", "F. Ingreso DPRH Planta/Servicios",
        "Asistencia", "Categoría Actual", "F. Alta Categoría", "Situación Revista", "Obs Revista",
        "Lugar de Trabajo", "Especif. Lugar Trabajo", "Sentencia", "Actividad", "Función Específica",
        "Título Profesional", "Dirección Organismo", "Área Trabajo", "Departamento", "División", "Sector",
        "Gremio Adherido", "Estado"
    ]
    ws_personal.append(headers_p)
    for col_idx in range(1, len(headers_p) + 1):
        cell = ws_personal.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    cursor.execute("""
        SELECT p.id, p.legajo, p.cuil, p.dni, p.apellido, p.nombre, p.fecha_nacimiento, p.genero,
               p.direccion, p.telefono, p.email, p.horario,
               p.fecha_ingreso_admin_pub, p.fecha_ingreso_recursos_hidricos_obra, p.fecha_ingreso_recursos_hidricos_planta,
               p.asistencia, p.categoria_actual, p.categoria_fecha_alta, p.situacion_revista, p.situacion_revista_obs,
               p.lugar_trabajo, p.lugar_trabajo_otro, p.sentencia, p.actividad, p.funcion,
               p.profesional_titulo, p.direccion_organismo, COALESCE(p.area_trabajo, a.nombre), p.departamento_trabajo, p.division_trabajo, p.sector_trabajo,
               p.gremio, p.estado
        FROM personal p
        LEFT JOIN areas a ON p.area_id = a.id
        WHERE p.deleted_at IS NULL
        ORDER BY p.apellido ASC, p.nombre ASC
    """)
    for row in cursor.fetchall():
        ws_personal.append(list(row))

    for row in ws_personal.iter_rows(min_row=2, max_row=ws_personal.max_row, min_col=1, max_col=len(headers_p)):
        for cell in row:
            cell.font = regular_font
            cell.border = thin_border

    # -------------------------------------------------------------
    # 3. PESTAÑA LICENCIAS
    # -------------------------------------------------------------
    ws_licencias = wb.create_sheet(title="Licencias")
    headers_l = ["ID", "DNI", "Agente", "Tipo Licencia", "Período Anual", "Fecha Inicio", "Fecha Fin", "Días Totales", "Familiar / Parentesco", "Diagnóstico / Motivo", "Resolución N°", "Médico", "Matrícula", "Estado Justificación"]
    ws_licencias.append(headers_l)
    for col_idx in range(1, len(headers_l) + 1):
        cell = ws_licencias.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    cursor.execute("""
        SELECT l.id, p.dni, (p.apellido || ', ' || p.nombre), l.tipo_licencia, l.periodo_anual,
               l.fecha_inicio, l.fecha_fin, l.dias_totales, 
               CASE WHEN l.familiar_nombre IS NOT NULL THEN (l.familiar_nombre || ' (' || COALESCE(l.familiar_parentesco,'') || ')') ELSE '' END,
               l.diagnostico, 
               l.certificado_numero, l.medico_nombre, l.medico_matricula, l.estado_justificacion
        FROM licencias l
        JOIN personal p ON l.personal_id = p.id
        WHERE l.deleted_at IS NULL
        ORDER BY l.fecha_inicio DESC
    """)
    for row in cursor.fetchall():
        ws_licencias.append(list(row))

    for row in ws_licencias.iter_rows(min_row=2, max_row=ws_licencias.max_row, min_col=1, max_col=len(headers_l)):
        for cell in row:
            cell.font = regular_font
            cell.border = thin_border

    # -------------------------------------------------------------
    # 4. PESTAÑA CONTRATOS
    # -------------------------------------------------------------
    ws_contratos = wb.create_sheet(title="Contratos")
    headers_c = ["ID", "DNI", "Agente", "Tipo Contrato", "Fecha Inicio", "Fecha Fin", "Monto Honorarios", "Resolución N°", "Estado"]
    ws_contratos.append(headers_c)
    for col_idx in range(1, len(headers_c) + 1):
        cell = ws_contratos.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    cursor.execute("""
        SELECT c.id, p.dni, (p.apellido || ', ' || p.nombre), c.tipo_contrato, 
               c.fecha_inicio, c.fecha_fin, c.monto_honorarios, c.resolucion_numero, c.estado
        FROM contratos c
        JOIN personal p ON c.personal_id = p.id
        ORDER BY c.fecha_inicio DESC
    """)
    for row in cursor.fetchall():
        ws_contratos.append(list(row))

    for row in ws_contratos.iter_rows(min_row=2, max_row=ws_contratos.max_row, min_col=1, max_col=len(headers_c)):
        for cell in row:
            cell.font = regular_font
            cell.border = thin_border

    # Ajustar anchos de columnas automáticamente
    for sheet in wb.worksheets:
        for col in sheet.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val = str(cell.value or '')
                if len(val) > max_len:
                    max_len = len(val)
            sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

    conn.close()

    # Guardar en memoria
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    registrar_auditoria("EXPORTAR_EXCEL", "sistema", None, "Exportación completa del libro Excel de 4 pestañas")
    return output

def inspeccionar_archivo_excel(file_bytes: bytes):
    """Devuelve las hojas disponibles y una vista previa de las primeras filas"""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sheets_info = []
    
    for name in wb.sheetnames:
        sheet = wb[name]
        preview_rows = []
        for row in sheet.iter_rows(values_only=True):
            if any(row):
                preview_rows.append([str(c) if c is not None else "" for c in row])
            if len(preview_rows) >= 6:
                break
        
        headers = preview_rows[0] if preview_rows else []
        data_preview = preview_rows[1:] if len(preview_rows) > 1 else []

        sheets_info.append({
            "nombre": name,
            "total_filas_aprox": sheet.max_row,
            "encabezados": headers,
            "vista_previa": data_preview
        })

    return sheets_info

def importar_personal_excel(file_bytes: bytes, sheet_name: str, modo: str = "upsert", usuario: str = "Admin"):
    """
    Importa empleados desde la hoja seleccionada.
    Unifica usando DNI como clave primaria (actualiza si existe o crea si es nuevo).
    """
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    if sheet_name not in wb.sheetnames:
        sheet = wb.active
    else:
        sheet = wb[sheet_name]

    conn = get_db_connection()
    cursor = conn.cursor()

    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        conn.close()
        return {"error": "La hoja seleccionada está vacía"}

    # Mapeo flexible de encabezados
    header_row = [str(c or '').strip().lower() for c in rows[0]]
    
    def find_col_idx(aliases):
        for alias in aliases:
            for idx, h in enumerate(header_row):
                if alias in h:
                    return idx
        return -1

    idx_dni = find_col_idx(["dni", "documento", "cedula", "identificacion"])
    idx_nombre = find_col_idx(["nombre", "nombres"])
    idx_apellido = find_col_idx(["apellido", "apellidos"])
    idx_legajo = find_col_idx(["legajo", "matricula", "nro_legajo"])
    idx_email = find_col_idx(["email", "correo", "mail"])
    idx_telefono = find_col_idx(["telefono", "tel", "celular", "movil"])
    idx_revista = find_col_idx(["revista", "situacion", "situacion_revista", "condicion", "planta"])
    idx_cargo = find_col_idx(["cargo", "puesto"])
    idx_funcion = find_col_idx(["funcion", "tarea", "rol"])
    idx_area = find_col_idx(["area", "sector", "departamento", "direccion"])
    idx_ingreso = find_col_idx(["ingreso", "fecha_ingreso", "alta"])

    if idx_dni == -1 or (idx_nombre == -1 and idx_apellido == -1):
        conn.close()
        return {"error": "No se encontraron columnas obligatorias (DNI y Nombre/Apellido). Verifique los encabezados."}

    # Cache de áreas
    cursor.execute("SELECT id, LOWER(nombre) FROM areas")
    areas_map = {row[1]: row[0] for row in cursor.fetchall()}

    nuevos = 0
    actualizados = 0
    errores = 0

    for row_idx, row in enumerate(rows[1:], start=2):
        if not any(row):
            continue
        
        dni_val = str(row[idx_dni] or '').strip().replace(".", "").replace(",", "").replace("-", "")
        if not dni_val or len(dni_val) < 4:
            continue

        nombre_val = str(row[idx_nombre] or '').strip() if idx_nombre != -1 and row[idx_nombre] else "Sin Nombre"
        apellido_val = str(row[idx_apellido] or '').strip() if idx_apellido != -1 and row[idx_apellido] else ""
        legajo_val = str(row[idx_legajo] or '').strip() if idx_legajo != -1 and row[idx_legajo] else None
        email_val = str(row[idx_email] or '').strip() if idx_email != -1 and row[idx_email] else None
        tel_val = str(row[idx_telefono] or '').strip() if idx_telefono != -1 and row[idx_telefono] else None
        revista_val = str(row[idx_revista] or '').strip() if idx_revista != -1 and row[idx_revista] else "Planta Permanente"
        cargo_val = str(row[idx_cargo] or '').strip() if idx_cargo != -1 and row[idx_cargo] else "Técnico"
        funcion_val = str(row[idx_funcion] or '').strip() if idx_funcion != -1 and row[idx_funcion] else "Operativo"
        ingreso_val = str(row[idx_ingreso] or '').strip() if idx_ingreso != -1 and row[idx_ingreso] else datetime.now().strftime("%Y-%m-%d")

        # Buscar o asociar área
        area_id = None
        if idx_area != -1 and row[idx_area]:
            area_str = str(row[idx_area]).strip().lower()
            if area_str in areas_map:
                area_id = areas_map[area_str]
            else:
                # Crear área si no existe
                cursor.execute("INSERT INTO areas (nombre, nivel) VALUES (?, 2)", (str(row[idx_area]).strip(),))
                area_id = cursor.lastrowid
                areas_map[area_str] = area_id

        # Verificar si DNI ya existe
        cursor.execute("SELECT id FROM personal WHERE dni = ?", (dni_val,))
        existente = cursor.fetchone()

        try:
            if existente:
                if modo in ["upsert", "actualizar"]:
                    cursor.execute("""
                        UPDATE personal SET
                            legajo = COALESCE(?, legajo),
                            nombre = ?,
                            apellido = ?,
                            email = COALESCE(?, email),
                            telefono = COALESCE(?, telefono),
                            situacion_revista = ?,
                            cargo = ?,
                            funcion = ?,
                            area_id = COALESCE(?, area_id),
                            deleted_at = NULL,
                            updated_at = CURRENT_TIMESTAMP
                        WHERE id = ?
                    """, (legajo_val, nombre_val, apellido_val, email_val, tel_val, revista_val, cargo_val, funcion_val, area_id, existente[0]))
                    actualizados += 1
            else:
                cursor.execute("""
                    INSERT INTO personal (
                        dni, legajo, nombre, apellido, email, telefono,
                        fecha_ingreso, situacion_revista, cargo, funcion,
                        area_id, estado, created_at, updated_at
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'Activo', CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
                """, (dni_val, legajo_val, nombre_val, apellido_val, email_val, tel_val, ingreso_val, revista_val, cargo_val, funcion_val, area_id))
                nuevos += 1
        except Exception as e:
            print(f"Error procesando fila {row_idx}: {e}")
            errores += 1

    conn.commit()
    conn.close()

    registrar_auditoria("IMPORTAR_EXCEL", "personal", None, f"Importación Excel: {nuevos} creados, {actualizados} actualizados, {errores} errores", usuario)

    return {
        "success": True,
        "nuevos": nuevos,
        "actualizados": actualizados,
        "errores": errores,
        "total_procesados": nuevos + actualizados + errores
    }

def exportar_personal_seleccion_excel(ids: list) -> bytes:
    """Genera un archivo Excel (.xlsx) con los datos del personal seleccionado por ID."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Personal Seleccionado"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    regular_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E0'),
        right=Side(style='thin', color='CBD5E0'),
        top=Side(style='thin', color='CBD5E0'),
        bottom=Side(style='thin', color='CBD5E0')
    )

    headers = [
        "ID", "Legajo", "CUIL", "DNI", "Apellido", "Nombre", "Fecha Nacimiento", "Sexo", 
        "Domicilio", "Teléfono", "Email", "Horario Trabajo", 
        "F. Ingreso Admin Pública", "F. Ingreso DPRH Obra", "F. Ingreso DPRH Planta/Servicios",
        "Asistencia", "Categoría Actual", "F. Alta Categoría", "Situación Revista", "Obs Revista",
        "Lugar de Trabajo", "Especif. Lugar Trabajo", "Sentencia", "Actividad", "Función Específica",
        "Título Profesional", "Dirección Organismo", "Área Asignada", "Departamento", "División", "Sector",
        "Gremio Adherido", "Estado"
    ]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    if ids:
        conn = get_db_connection()
        cursor = conn.cursor()
        placeholders = ",".join(["?"] * len(ids))
        cursor.execute(f"""
            SELECT p.id, p.legajo, p.cuil, p.dni, p.apellido, p.nombre, p.fecha_nacimiento, p.genero,
                   p.direccion, p.telefono, p.email, p.horario,
                   p.fecha_ingreso_admin_pub, p.fecha_ingreso_recursos_hidricos_obra, p.fecha_ingreso_recursos_hidricos_planta,
                   p.asistencia, p.categoria_actual, p.categoria_fecha_alta, p.situacion_revista, p.situacion_revista_obs,
                   p.lugar_trabajo, p.lugar_trabajo_otro, p.sentencia, p.actividad, p.funcion,
                   p.profesional_titulo, p.direccion_organismo, COALESCE(p.area_trabajo, a.nombre), p.departamento_trabajo, p.division_trabajo, p.sector_trabajo,
                   p.gremio, p.estado
            FROM personal p
            LEFT JOIN areas a ON p.area_id = a.id
            WHERE p.id IN ({placeholders}) AND p.deleted_at IS NULL
            ORDER BY p.apellido ASC, p.nombre ASC
        """, ids)
        for row in cursor.fetchall():
            ws.append(list(row))
        conn.close()

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.font = regular_font
            cell.border = thin_border

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

def generar_plantilla_ejemplo():
    """Genera una plantilla Excel (.xlsx) con columnas oficiales requeridas y filas de ejemplo."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Personal"
    ws.views.sheetView[0].showGridLines = True

    header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    regular_font = Font(name="Arial", size=10)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E0'),
        right=Side(style='thin', color='CBD5E0'),
        top=Side(style='thin', color='CBD5E0'),
        bottom=Side(style='thin', color='CBD5E0')
    )

    headers = [
        "DNI", "Nombres", "Apellidos", "Legajo", "Situacion Revista",
        "Area", "Cargo", "Funcion", "Lugar de Trabajo",
        "Fecha Ingreso", "Email", "Telefono", "Direccion"
    ]
    ws.append(headers)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ejemplos = [
        ["28123456", "Carlos Alberto", "Rodríguez", "LEG-1001", "Planta Permanente", "Área de Recursos Humanos", "Inspector Técnico", "Inspector", "Oficina Central", "2015-03-01", "crodriguez@hidricos.gob.ar", "3884123456", "Av. Fascio 750, S.S. de Jujuy"],
        ["33987654", "Mariana Beatriz", "González", "LEG-1042", "Contrato de Locación de Servicios", "División de Obras", "Supervisora de Obras", "Profesional", "Las Maderas", "2020-06-15", "mgonzalez@hidricos.gob.ar", "3885987654", "Barrio Los Perales, S.S. de Jujuy"],
        ["35456789", "Jorge David", "Flores", "LEG-1088", "Planta Permanente", "Departamento Manejo Integrado de Diques", "Operador de Válvulas", "Guardadique", "La Ciénaga", "2018-09-01", "jflores@hidricos.gob.ar", "3884888999", "El Carmen, Jujuy"]
    ]

    for fila in ejemplos:
        ws.append(fila)

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
        for cell in row:
            cell.font = regular_font
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center")

    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

generar_plantilla_excel_ejemplo = generar_plantilla_ejemplo


