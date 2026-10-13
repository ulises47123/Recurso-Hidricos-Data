"""
tests/test_catalogos_postgres.py
Suite de pruebas de integración y unitarias para:
  1. Nuevos endpoints de Catálogos dinámicos (/api/catalogos, /api/catalogos/situacion_revista, /api/catalogos/normativa).
  2. Descarga y formato oficial de Plantilla Excel de ejemplo (/api/excel/plantilla-ejemplo).
  3. Edición de contraseña de administrador a 'admin123' y verificación de autenticación.
  4. Compatibilidad del script de migración SQLite a PostgreSQL (conversión de tipos, DDL, secuencias y triggers).
"""

import os
import sys
import io
import json
import asyncio
import unittest
from unittest.mock import MagicMock, patch

ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT_DIR not in sys.path:
    sys.path.insert(0, ROOT_DIR)

import openpyxl
from fastapi import HTTPException
from backend.database import init_db, get_db_connection
from backend.models import UsuarioLogin, UsuarioUpdate
from backend.main import (
    obtener_catalogos,
    actualizar_catalogo_situacion_revista,
    actualizar_catalogo_normativa,
    descargar_plantilla_ejemplo,
    login,
    actualizar_usuario,
    listar_usuarios
)
from backend.services import excel_service
from backend.migrate_to_postgres import (
    convertir_tipo_sqlite_a_postgres,
    generar_ddl_tabla_dinamica,
    obtener_tablas_sqlite,
    obtener_columnas_sqlite,
    reajustar_secuencias_postgres,
    migrar_sqlite_a_postgres,
    POSTGRES_SCHEMA_DDL,
    POSTGRES_INDEXES_DDL,
    TABLES_TOPOLOGICAL_ORDER
)


class TestCatalogosPostgres(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        init_db()
        from backend.services import organigrama_service
        areas = organigrama_service.listar_areas()
        if not areas:
            organigrama_service.crear_area({
                "nombre": "Dirección Provincial de Recursos Hídricos",
                "codigo": "DPRH-01",
                "nivel": 1
            })

    # =========================================================================
    # 1. PRUEBAS DE ENDPOINTS DE CATÁLOGOS DINÁMICOS
    # =========================================================================
    def test_obtener_catalogos_estructura_completa(self):
        """Valida que /api/catalogos retorne todas las claves requeridas y no vacías."""
        cat = obtener_catalogos()
        self.assertIsInstance(cat, dict)

        claves_esperadas = [
            "situaciones_revista",
            "tipos_licencia",
            "parentescos",
            "tipos_documento",
            "tipos_contrato",
            "estados_personal",
            "estados_justificacion",
            "areas"
        ]
        for clave in claves_esperadas:
            self.assertIn(clave, cat, f"Falta la clave '{clave}' en la respuesta de catálogos")
            self.assertIsInstance(cat[clave], list, f"'{clave}' debe ser una lista")
            self.assertGreater(len(cat[clave]), 0, f"'{clave}' no debe estar vacía")

    def test_actualizar_catalogo_situacion_revista(self):
        """Valida la actualización dinámica del catálogo de Situación de Revista."""
        nueva_lista = [
            "Planta Permanente",
            "Contratado Ley 5473",
            "Locación de Obras Especiales",
            "Gabinete de Ministros"
        ]
        res = actualizar_catalogo_situacion_revista(nueva_lista)
        self.assertTrue(res.get("success"), "El guardado del catálogo debe ser exitoso")

        # Verificar que al consultar /api/catalogos se devuelva la lista actualizada
        cat = obtener_catalogos()
        self.assertEqual(cat["situaciones_revista"], nueva_lista)
        self.assertIn("Locación de Obras Especiales", cat["situaciones_revista"])

        # Restaurar catálogo estándar
        lista_default = [
            "Planta Permanente",
            "Contratado Ley",
            "Locación de Servicios",
            "Contrato de Obra",
            "Gabinete"
        ]
        actualizar_catalogo_situacion_revista(lista_default)

    def test_actualizar_catalogo_normativa(self):
        """Valida la actualización de normativas dinámicas de licencias."""
        nueva_normativa = [
            {"tipo": "Médica Ordinaria", "dias_max": 30, "articulo": "Art. 10"},
            {"tipo": "Licencia Excepcional Hídrica", "dias_max": 15, "articulo": "Art. 99 DPRH"}
        ]
        res = actualizar_catalogo_normativa(nueva_normativa)
        self.assertTrue(res.get("success"))

        conn = get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT valor FROM configuracion WHERE clave = 'catalogo_normativa'")
        row = cursor.fetchone()
        conn.close()

        self.assertIsNotNone(row)
        normativas_guardadas = json.loads(row[0])
        tipos = [n["tipo"] for n in normativas_guardadas]
        self.assertIn("Licencia Excepcional Hídrica", tipos)

    # =========================================================================
    # 2. PRUEBAS DE DESCARGA DE PLANTILLA EXCEL DE EJEMPLO
    # =========================================================================
    def test_descarga_plantilla_ejemplo_endpoint(self):
        """Valida la descarga de la plantilla Excel oficial mediante el endpoint."""
        resp = descargar_plantilla_ejemplo()
        self.assertEqual(resp.status_code, 200)
        self.assertIn("Plantilla_Ejemplo_Personal_DPRH.xlsx", resp.headers.get("Content-Disposition", ""))

        # Leer el contenido binario del streaming response de forma asíncrona
        async def _leer_stream():
            chunks = []
            async for chunk in resp.body_iterator:
                chunks.append(chunk)
            return b"".join(chunks)

        excel_bytes = asyncio.run(_leer_stream())
        self.assertGreater(len(excel_bytes), 0, "El archivo Excel no debe estar vacío")
        self.assertTrue(excel_bytes.startswith(b"PK"), "El archivo debe ser un formato OpenXML (.xlsx)")

        # Cargar el libro Excel en memoria y verificar su estructura
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertIn("Personal", wb.sheetnames, "La plantilla debe contener la hoja 'Personal'")

        ws = wb["Personal"]
        columnas_obtenidas = [ws.cell(row=1, column=col).value for col in range(1, 14)]
        columnas_esperadas = [
            "DNI", "Nombres", "Apellidos", "Legajo", "Situacion Revista",
            "Area", "Cargo", "Funcion", "Lugar de Trabajo",
            "Fecha Ingreso", "Email", "Telefono", "Direccion"
        ]
        self.assertEqual(columnas_obtenidas, columnas_esperadas)

        # Debe contener al menos 3 filas de agentes de ejemplo precargados
        filas_datos = list(ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True))
        filas_validas = [f for f in filas_datos if f[0] is not None]
        self.assertGreaterEqual(len(filas_validas), 3, "Debe haber al menos 3 filas de ejemplo")
        self.assertIn("Carlos Alberto", [f[1] for f in filas_validas])

    def test_generar_plantilla_ejemplo_service(self):
        """Valida directamente el servicio generador de plantillas Excel."""
        buf = excel_service.generar_plantilla_ejemplo()
        self.assertIsInstance(buf, io.BytesIO)
        contenido = buf.getvalue()
        self.assertGreater(len(contenido), 1000)
        self.assertTrue(contenido.startswith(b"PK"))

    # =========================================================================
    # 3. PRUEBAS DE EDICIÓN DE CONTRASEÑA DE ADMIN ('admin123')
    # =========================================================================
    def test_cambio_password_admin_y_autenticacion(self):
        """Valida la actualización de contraseña del usuario admin a 'admin123' y el login."""
        usuarios = listar_usuarios()
        admin_user = next((u for u in usuarios if u["username"] == "admin"), None)
        self.assertIsNotNone(admin_user, "El usuario 'admin' debe existir")

        admin_id = admin_user["id"]

        # 1. Cambiar la contraseña a 'admin123'
        res_update = actualizar_usuario(admin_id, UsuarioUpdate(password="admin123"))
        self.assertTrue(res_update.get("success"), "La contraseña de admin debe actualizarse")

        # 2. Iniciar sesión con la nueva contraseña 'admin123' (debe ser exitoso)
        login_res = login(UsuarioLogin(username="admin", password="admin123"))
        self.assertTrue(login_res.get("success"))
        self.assertEqual(login_res["usuario"]["username"], "admin")
        self.assertEqual(login_res["usuario"]["rol"], "admin")

        # 3. Iniciar sesión con contraseña incorrecta (debe fallar con HTTP 401)
        with self.assertRaises(HTTPException) as cm:
            login(UsuarioLogin(username="admin", password="password_incorrecta"))
        self.assertEqual(cm.exception.status_code, 401)

        # 4. Restaurar la contraseña a 'admin' para mantener idempotencia
        actualizar_usuario(admin_id, UsuarioUpdate(password="admin"))
        login_revert = login(UsuarioLogin(username="admin", password="admin"))
        self.assertTrue(login_revert.get("success"))

    # =========================================================================
    # 4. PRUEBAS DE COMPATIBILIDAD DE MIGRACIÓN SQLITE A POSTGRESQL
    # =========================================================================
    def test_conversion_tipos_datos_sqlite_a_postgres(self):
        """Valida la conversión exacta de tipos de datos para PostgreSQL."""
        self.assertEqual(
            convertir_tipo_sqlite_a_postgres("id", "INTEGER", pk=True, autoincrement=True),
            "SERIAL PRIMARY KEY"
        )
        self.assertEqual(
            convertir_tipo_sqlite_a_postgres("monto_honorarios", "REAL"),
            "DOUBLE PRECISION"
        )
        self.assertEqual(
            convertir_tipo_sqlite_a_postgres("nombre", "TEXT"),
            "TEXT"
        )
        self.assertEqual(
            convertir_tipo_sqlite_a_postgres("activo", "INTEGER"),
            "INTEGER"
        )
        self.assertEqual(
            convertir_tipo_sqlite_a_postgres("archivo_bin", "BLOB"),
            "BYTEA"
        )

    def test_esquema_ddl_completo_para_todas_las_tablas(self):
        """Valida que todas las tablas del sistema cuenten con DDL PostgreSQL oficial."""
        tablas_requeridas = [
            "areas", "personal", "licencias", "contratos", "familiares",
            "historial_funciones", "documentos_legajo", "usuarios",
            "adicionales_personal", "historial_categorias", "beneficios_articulo_79",
            "suspensiones_licencia", "auditoria_logs", "configuracion",
            "compensatorios_saldo", "compensatorios_cargas", "compensatorios_solicitudes",
            "feriados"
        ]
        for t in tablas_requeridas:
            self.assertIn(t, POSTGRES_SCHEMA_DDL, f"Falta DDL para la tabla '{t}'")
            ddl = POSTGRES_SCHEMA_DDL[t]
            self.assertIn(f"CREATE TABLE IF NOT EXISTS {t}", ddl)

    def test_ddl_dinamica_tablas_extras(self):
        """Valida la generación dinámica de DDL para tablas no predefinidas."""
        import sqlite3
        conn = sqlite3.connect(":memory:")
        cur = conn.cursor()
        cur.execute("""
            CREATE TABLE mi_tabla_custom (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                titulo TEXT NOT NULL,
                valor REAL DEFAULT 0.0
            )
        """)
        ddl = generar_ddl_tabla_dinamica(conn, "mi_tabla_custom")
        self.assertIn("CREATE TABLE IF NOT EXISTS mi_tabla_custom", ddl)
        self.assertIn('"id" SERIAL PRIMARY KEY', ddl)
        self.assertIn('"titulo" TEXT NOT NULL', ddl)
        self.assertIn('"valor" DOUBLE PRECISION', ddl)
        conn.close()

    def test_reajuste_secuencias_postgres_query(self):
        """Valida la generación de sentencias setval() para secuencias PostgreSQL."""
        mock_cur = MagicMock()
        mock_cur.fetchall.return_value = [("personal", "id", "nextval('personal_id_seq'::regclass)")]
        mock_cur.fetchone.side_effect = [("personal_id_seq",), (42,)]

        ajustes = reajustar_secuencias_postgres(mock_cur)
        self.assertEqual(len(ajustes), 1)
        self.assertIn("personal.id -> secuencia 'personal_id_seq' en 42", ajustes[0])

    def test_simulacion_migracion_completa_mock(self):
        """
        Simula el flujo completo de migrar_sqlite_a_postgres usando mocks de conexión
        para verificar desactivación de triggers, DDL, inserción por lotes y setval.
        """
        mock_pg_conn = MagicMock()
        mock_pg_cur = MagicMock()
        mock_pg_conn.cursor.return_value = mock_pg_cur

        mock_pg_cur.fetchall.return_value = []

        with patch("psycopg2.connect", return_value=mock_pg_conn), \
             patch("psycopg2.extras.execute_batch") as mock_exec_batch:

            resumen = migrar_sqlite_a_postgres(
                sqlite_path=os.path.join(ROOT_DIR, "recursos_hidricos.db"),
                host="localhost",
                port=5432,
                user="postgres",
                password="secret_password",
                dbname="recursos_hidricos",
                drop_existing=False,
                verbose=False
            )

            self.assertTrue(resumen["success"])
            self.assertGreater(len(resumen["tablas_migradas"]), 0)

            # Verificar que se intentó desactivar y luego reactivar foreign keys y triggers
            sentencias_ejecutadas = [str(call[0][0]) for call in mock_pg_cur.execute.call_args_list if call[0]]
            self.assertTrue(
                any("SET session_replication_role = 'replica'" in s for s in sentencias_ejecutadas),
                "Debe desactivar claves foráneas con replica role"
            )
            self.assertTrue(
                any("SET session_replication_role = 'origin'" in s for s in sentencias_ejecutadas),
                "Debe reactivar claves foráneas con origin role"
            )

            # Verificar que se crearon los índices
            self.assertTrue(
                any("CREATE INDEX IF NOT EXISTS idx_personal_dni" in s for s in sentencias_ejecutadas),
                "Debe crear los índices en PostgreSQL"
            )


if __name__ == "__main__":
    unittest.main()
